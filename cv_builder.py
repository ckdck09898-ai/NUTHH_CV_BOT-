import os, re, textwrap
from pathlib import Path
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Image as RLImage, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from templates import get_template

ROOT = Path(__file__).parent
OUT = ROOT / "generated_cvs"
PREVIEWS = ROOT / "previews"
FONTS = ROOT / "fonts"
OUT.mkdir(exist_ok=True); PREVIEWS.mkdir(exist_ok=True)

FIELDS = [
    ("name", "Full name / ឈ្មោះពេញ", True),
    ("job_title", "Target job / មុខតំណែងដែលចង់ដាក់", True),
    ("phone", "Phone / លេខទូរស័ព្ទ", True),
    ("email", "Email / អ៊ីមែល", False),
    ("address", "Address / អាសយដ្ឋាន", False),
    ("summary", "Career objective / គោលបំណងការងារ", False),
    ("education", "Education / ការសិក្សា", False),
    ("skills", "Skills / ជំនាញ (separate with commas)", False),
    ("experience", "Experience / បទពិសោធន៍ (write None if none)", False),
    ("projects", "Projects / គម្រោង", False),
    ("languages", "Languages / ភាសា", False),
    ("certificates", "Certificates / វិញ្ញាបនបត្រ", False),
]
SECTION_LABELS = {
    "en": {"job_title":"TARGET POSITION", "summary":"PROFILE / OBJECTIVE", "education":"EDUCATION", "skills":"SKILLS", "experience":"EXPERIENCE", "projects":"PROJECTS", "languages":"LANGUAGES", "certificates":"CERTIFICATES", "contact":"CONTACT"},
    "km": {"job_title":"មុខតំណែងដែលចង់ដាក់", "summary":"គោលបំណងការងារ", "education":"ការសិក្សា", "skills":"ជំនាញ", "experience":"បទពិសោធន៍", "projects":"គម្រោង", "languages":"ភាសា", "certificates":"វិញ្ញាបនបត្រ", "contact":"ទំនាក់ទំនង"},
    "zh": {"job_title":"求职职位", "summary":"个人简介", "education":"教育背景", "skills":"技能", "experience":"工作经验", "projects":"项目", "languages":"语言", "certificates":"证书", "contact":"联系方式"},
    "fr": {"job_title":"POSTE VISÉ", "summary":"PROFIL", "education":"FORMATION", "skills":"COMPÉTENCES", "experience":"EXPÉRIENCE", "projects":"PROJETS", "languages":"LANGUES", "certificates":"CERTIFICATS", "contact":"CONTACT"},
}

def _hex(hex_color):
    return colors.HexColor(hex_color)

def _register_font():
    candidates = [FONTS/"NotoSansKhmer-Regular.ttf", FONTS/"NotoSansKhmer.ttf", ROOT/"NotoSansKhmer-Regular.ttf"]
    for path in candidates:
        if path.exists():
            try:
                if "NotoKhmerCV" not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont("NotoKhmerCV", str(path)))
                return "NotoKhmerCV"
            except Exception:
                pass
    return "Helvetica"

def make_preview(template_id, output_path=None, data=None):
    t = get_template(template_id)
    output_path = Path(output_path or PREVIEWS / f"{template_id}.png")
    w, h = 720, 1000
    img = Image.new("RGB", (w,h), "white")
    d = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("DejaVuSans.ttf", 32)
        body_font = ImageFont.truetype("DejaVuSans.ttf", 17)
        small_font = ImageFont.truetype("DejaVuSans.ttf", 14)
    except OSError:
        title_font = body_font = small_font = ImageFont.load_default()
    color = t["color"]
    if t["layout"] == "Sidebar":
        d.rectangle((0,0,220,h), fill=color)
        left = 255
        d.ellipse((58,55,162,159), fill="#FFFFFF")
        d.text((40,185), "CONTACT", fill="white", font=body_font)
        for y in (225,255,285,315): d.rounded_rectangle((38,y,180,y+8), radius=4, fill="#DCE5F5")
    else:
        d.rectangle((0,0,w,150), fill=color)
        left = 52
    display_name = ((data or {}).get("name") or "YOUR NAME")[:28]
    display_job = ((data or {}).get("job_title") or "TARGET POSITION")[:36]
    d.text((left, 52 if t["layout"] != "Sidebar" else 55), display_name, font=title_font, fill=("white" if t["layout"] != "Sidebar" else color))
    d.text((left, 105 if t["layout"] != "Sidebar" else 105), display_job, font=small_font, fill=("#E8EEFF" if t["layout"] != "Sidebar" else "#64748B"))
    y = 190 if t["layout"] != "Sidebar" else 210
    for title in ("PROFILE", "EDUCATION", "SKILLS", "EXPERIENCE", "PROJECTS"):
        d.text((left,y), title, font=body_font, fill=color)
        y += 35
        for j in range(2):
            d.rounded_rectangle((left,y,left+min(390, 470-j*75),y+8), radius=4, fill="#D8DEE8")
            y += 22
        y += 22
    d.rectangle((20,20,w-20,h-20), outline="#E4E7EC", width=2)
    d.text((w-190,h-48), f"TEMPLATE {t['number']:03d}", fill="#98A2B3", font=small_font)
    img.save(output_path)
    return str(output_path)

def build_pdf(data, template_id, language="en", photo_path=None, user_id="user"):
    t = get_template(template_id)
    path = OUT / f"cv_{user_id}_{template_id}.pdf"
    font = _register_font()
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=38, bottomMargin=38)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CVTitle", parent=styles["Title"], fontName=font, fontSize=23, leading=29, textColor=_hex(t["color"]), alignment=TA_CENTER, spaceAfter=5)
    subtitle_style = ParagraphStyle("CVSub", parent=styles["Normal"], fontName=font, fontSize=10, leading=15, textColor=colors.HexColor("#475467"), alignment=TA_CENTER, spaceAfter=6)
    head_style = ParagraphStyle("CVHead", parent=styles["Heading2"], fontName=font, fontSize=11, leading=16, textColor=_hex(t["color"]), spaceBefore=11, spaceAfter=4)
    body_style = ParagraphStyle("CVBody", parent=styles["BodyText"], fontName=font, fontSize=9.5, leading=14, textColor=colors.HexColor("#273142"), spaceAfter=4, wordWrap="CJK")
    labels = SECTION_LABELS.get(language, SECTION_LABELS["en"])
    story=[]
    # A photo is optional; failure to decode it should not prevent CV generation.
    if photo_path and Path(photo_path).exists():
        try:
            photo = RLImage(photo_path, width=72, height=72)
            header = Table([[photo, Paragraph(escape(data.get("name", "YOUR NAME")), title_style)]], colWidths=[82, 420])
            header.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
            story.append(header)
        except Exception:
            story.append(Paragraph(escape(data.get("name", "YOUR NAME")), title_style))
    else:
        story.append(Paragraph(escape(data.get("name", "YOUR NAME")), title_style))
    if data.get("job_title"):
        story.append(Paragraph(escape(data["job_title"]), subtitle_style))
    contact = "  |  ".join(data.get(k, "").strip() for k in ("phone","email","address") if data.get(k, "").strip())
    if contact: story.append(Paragraph(escape(contact), subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.3, color=_hex(t["color"]), spaceBefore=5, spaceAfter=8))
    for key in ("summary","education","skills","experience","projects","languages","certificates"):
        value = str(data.get(key, "")).strip()
        if not value: continue
        story.append(Paragraph(escape(labels.get(key, key.upper())), head_style))
        for line in value.splitlines():
            if line.strip(): story.append(Paragraph(escape(line.strip()), body_style))
    doc.build(story)
    return str(path)

def build_docx(data, template_id, language="en", photo_path=None, user_id="user"):
    t = get_template(template_id)
    path = OUT / f"cv_{user_id}_{template_id}.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.55); section.bottom_margin = Inches(.55)
    p = doc.add_paragraph(); p.alignment = 1
    r = p.add_run(data.get("name", "YOUR NAME")); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor.from_string(t["color"].replace("#", ""))
    p = doc.add_paragraph(); p.alignment=1
    p.add_run(data.get("job_title", ""))
    contact = " | ".join(data.get(k, "").strip() for k in ("phone","email","address") if data.get(k, "").strip())
    if contact:
        p=doc.add_paragraph(contact); p.alignment=1
    if photo_path and Path(photo_path).exists():
        try: doc.add_picture(photo_path, width=Inches(1.0))
        except Exception: pass
    labels = SECTION_LABELS.get(language, SECTION_LABELS["en"])
    for key in ("summary","education","skills","experience","projects","languages","certificates"):
        value=str(data.get(key, "")).strip()
        if not value: continue
        h=doc.add_heading(labels.get(key,key.upper()), level=2)
        for run in h.runs:
            run.font.color.rgb=RGBColor.from_string(t["color"].replace("#", ""))
        for line in value.splitlines():
            if line.strip(): doc.add_paragraph(line.strip())
    doc.save(path)
    return str(path)

def ats_check(data, job_description=""):
    checks = {
        "Name included": bool(data.get("name", "").strip()),
        "Target job included": bool(data.get("job_title", "").strip()),
        "Contact information": bool(data.get("phone", "").strip() or data.get("email", "").strip()),
        "Career summary": bool(data.get("summary", "").strip()),
        "Education": bool(data.get("education", "").strip()),
        "Skills": bool(data.get("skills", "").strip()),
        "Experience or relevant projects": bool(data.get("experience", "").strip() or data.get("projects", "").strip()),
        "Languages": bool(data.get("languages", "").strip()),
    }
    score = round(sum(checks.values()) / len(checks) * 100)
    suggestions = [name for name, passed in checks.items() if not passed]
    match = None
    if job_description.strip():
        cv_text = " ".join(str(v) for k, v in data.items() if not str(k).startswith("_")).lower()
        words = {w for w in re.findall(r"[a-zA-Z]{3,}", job_description.lower()) if w not in {"and","the","for","with","you","are","our","will","from","that","this","job","work"}}
        hits = sorted(w for w in words if w in cv_text)
        match = {"score": round(len(hits)/max(1,len(words))*100), "matched": hits[:30], "missing": sorted(words-set(hits))[:30]}
    return {"score":score,"checks":checks,"suggestions":suggestions,"match":match}

def write_assistance(kind, data, language="en", job_description=""):
    """Free rule/template-based writing helper; not a generative AI model."""
    name = data.get("name", "[Your Name]").strip() or "[Your Name]"
    job = data.get("job_title", "the position").strip() or "the position"
    skills = [s.strip() for s in data.get("skills", "").replace("\n", ",").split(",") if s.strip()]
    skill_text = ", ".join(skills[:5]) if skills else "a willingness to learn, teamwork, and communication"
    no_exp = data.get("experience", "").strip().lower() in ("none", "no experience", "n/a", "មិនមាន", "អត់មាន", "")
    if language == "km":
        if kind == "summary":
            return f"ខ្ញុំមានបំណងដាក់ពាក្យសម្រាប់មុខតំណែង {job}។ ខ្ញុំមានជំនាញ {skill_text} មានការទទួលខុសត្រូវ និងមានឆន្ទៈរៀនសូត្របន្ថែម ដើម្បីរួមចំណែកដល់ក្រុមការងារ។" if no_exp else f"ខ្ញុំចាប់អារម្មណ៍លើមុខតំណែង {job} និងចង់ប្រើប្រាស់ជំនាញ {skill_text} ព្រមទាំងបទពិសោធន៍ដែលខ្ញុំមាន ដើម្បីរួមចំណែកដល់ក្រុមការងារ។"
        if kind == "cover":
            return f"ជូនចំពោះអ្នកទទួលបន្ទុកជ្រើសរើសបុគ្គលិក\n\nខ្ញុំសូមដាក់ពាក្យសម្រាប់មុខតំណែង {job}។ ខ្ញុំមានជំនាញ {skill_text} និងមានឆន្ទៈរៀនសូត្រ អនុវត្តតាមការណែនាំ និងធ្វើការជាក្រុម។ {('ទោះបីខ្ញុំមិនទាន់មានបទពិសោធន៍ការងារផ្លូវការក៏ដោយ ខ្ញុំត្រៀមខ្លួនរៀន និងអភិវឌ្ឍសមត្ថភាព។' if no_exp else 'ខ្ញុំចង់យកបទពិសោធន៍ និងជំនាញរបស់ខ្ញុំមកចូលរួមជាមួយក្រុមហ៊ុន។')}\n\nសូមអរគុណចំពោះការពិចារណាពាក្យសុំរបស់ខ្ញុំ។\n\nដោយក្តីគោរព\n{name}"
        return f"ប្រធានបទ៖ ដាក់ពាក្យសម្រាប់មុខតំណែង {job}\n\nជម្រាបសួរ ក្រុមការងារជ្រើសរើសបុគ្គលិក,\n\nខ្ញុំសូមផ្ញើ CV របស់ខ្ញុំសម្រាប់ការពិចារណាលើមុខតំណែង {job}។ សូមមេត្តាពិនិត្យឯកសារភ្ជាប់។\n\nសូមអរគុណ,\n{name}"
    if language == "zh":
        if kind == "summary": return f"我希望申请{job}职位。我具备{skill_text}，认真负责，并愿意持续学习和提升能力，为团队作出贡献。"
        if kind == "cover": return f"尊敬的招聘经理：\n\n我申请贵公司的{job}职位。我具备{skill_text}，并愿意学习、遵循工作流程和积极参与团队合作。感谢您考虑我的申请。\n\n此致\n{name}"
        return f"主题：申请{job}职位\n\n您好，招聘团队：\n\n随信附上我的简历，申请贵公司的{job}职位。感谢您的时间与考虑。\n\n此致\n{name}"
    if kind == "summary":
        intro = "Motivated entry-level candidate" if no_exp else "Motivated professional"
        return f"{intro} seeking a {job} position. Brings strengths in {skill_text}, a responsible work ethic, and a willingness to learn, contribute to team goals, and develop professionally."
    if kind == "cover":
        exp = "Although I am at the beginning of my career, I am eager to learn, follow procedures, and contribute reliably." if no_exp else "I am ready to apply my relevant experience and skills to support your team."
        return f"Dear Hiring Manager,\n\nI am writing to apply for the {job} position. My strengths include {skill_text}. {exp} I would welcome the opportunity to discuss how I can contribute to your organization.\n\nThank you for your time and consideration.\n\nSincerely,\n{name}"
    return f"Subject: Application for {job}\n\nDear Hiring Team,\n\nPlease find attached my CV for the {job} position. I would appreciate your consideration and welcome the opportunity to discuss my suitability for the role.\n\nThank you for your time.\n\nKind regards,\n{name}"
